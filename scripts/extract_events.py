import json, re
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

SOURCES = [
    (Path('/tmp/calendar-source/01.xlsx'), 'ฝ่ายบริหารงานทั่วไป', 'general'),
    (Path('/tmp/calendar-source/02.xlsx'), 'ฝ่ายบริหารงานบุคคลและกิจการนักเรียน', 'personnel'),
    (Path('/tmp/calendar-source/03.xlsx'), 'ฝ่ายบริหารงานวิชาการ', 'academic'),
]

MONTHS = {'ม.ค.':1,'ก.พ.':2,'มี.ค.':3,'เม.ย.':4,'พ.ค.':5,'มิ.ย.':6,'ก.ค.':7,'ส.ค.':8,'ก.ย.':9,'ต.ค.':10,'พ.ย.':11,'ธ.ค.':12}

def academic_year(ws):
    m = re.search(r'(25\d{2})', str(ws.cell(1,1).value or ''))
    if m: return int(m.group(1))
    m = re.search(r'(25\d{2})', ws.title)
    return int(m.group(1)) if m else None

def semester(ws):
    title = str(ws.cell(1,1).value or '')
    m = re.search(r'ภาคเรียนที่\s*([12])', title)
    if m: return int(m.group(1))
    return 2 if ws.title.startswith('2') else 1

def date_text(v, ay):
    if isinstance(v, datetime):
        by = v.year if v.year >= 2500 else (2500 + v.year % 100 if v.year >= 1900 else v.year + 543)
        return f'{v.day} {list(MONTHS)[v.month-1]} {str(by)[-2:]}'
    return re.sub(r'\s+', ' ', str(v or '')).strip()

def parse_date_value(v, ay, sem):
    if isinstance(v, datetime):
        by = v.year if v.year >= 2500 else (2500 + v.year % 100 if v.year >= 1900 else v.year + 543)
        try: return datetime(by-543, v.month, v.day)
        except ValueError: return None
    s = re.sub(r'\s+', ' ', str(v or '').strip()).replace('–','-')
    if not s: return None
    md = re.search(r'(\d{1,2})(?:\s*-\s*\d{1,2})?\s*(ม\.ค\.|ก\.พ\.|มี\.ค\.|เม\.ย\.|พ\.ค\.|มิ\.ย\.|ก\.ค\.|ส\.ค\.|ก\.ย\.|ต\.ค\.|พ\.ย\.|ธ\.ค\.)\s*(\d{2,4})?', s)
    if not md: return None
    day, mon, yr = int(md.group(1)), MONTHS[md.group(2)], md.group(3)
    if yr:
        by = int(yr)
        if by < 100: by += 2500
    else:
        by = ay + (1 if sem == 2 and mon <= 4 else 0)
    try: return datetime(by-543, mon, day)
    except ValueError: return None

events=[]
seen=set()
for path, division, division_key in SOURCES:
    wb=load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        ay=academic_year(ws)
        sem=semester(ws)
        if not ay: continue
        for row_idx,row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx <= 3: continue
            vals=list(row)
            for c in range(0,len(vals),3):
                if c+1 >= len(vals): continue
                dv, activity = vals[c], vals[c+1]
                owner = vals[c+2] if c+2 < len(vals) else ''
                if dv in (None,'') or activity in (None,''): continue
                activity=re.sub(r'\s+',' ',str(activity)).strip()
                owner=re.sub(r'\s+',' ',str(owner or 'ไม่ระบุ')).strip()
                dtext=date_text(dv, ay)
                start=parse_date_value(dv, ay, sem)
                key=(division,ay,sem,dtext,activity,owner)
                if key in seen: continue
                seen.add(key)
                events.append({
                    'id':f'e{len(events)+1}', 'academicYear':ay, 'semester':sem,
                    'dateText':dtext, 'start':start.strftime('%Y-%m-%d') if start else None,
                    'activity':activity, 'owner':owner,
                    'division':division, 'divisionKey':division_key,
                    'sourceSheet':ws.title
                })

events.sort(key=lambda e:(e['start'] or '9999-12-31',e['division'],e['activity']))
out=Path('/workspace/sites/bj-school-calendar/dist')
out.mkdir(parents=True,exist_ok=True)
(out/'events.json').write_text(json.dumps(events,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print(json.dumps({'events':len(events),'years':sorted({e['academicYear'] for e in events}), 'byDivision':{k:sum(e['divisionKey']==k for e in events) for k in ['general','personnel','academic']},'undated':sum(not e['start'] for e in events)},ensure_ascii=False))
